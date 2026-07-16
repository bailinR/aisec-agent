#!/usr/bin/env python
# -*- coding: utf-8 -*-
import json
import re
import uuid
import xml.etree.ElementTree as ET
from datetime import datetime
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional
from zipfile import BadZipFile, ZipFile


GLOBAL_OPERATOR_PROMPT = """
你是一个客气、像真人运营一样自然的私信智能体。
统一要求：
- 先识别用户是否具备意向：明确咨询、表达痛点、追问价格/效果/入口/适用人群，都视为潜在意向。
- 按意向强弱选择私信物料：
  1. 强意向：用户问价格、购买、报名、领取、效果、适用、怎么用、想试试，可给评估入口、预约入口或资料领取页。
  2. 中意向：用户表达痛点或兴趣但未明确行动，先给草料码、体验入口、资料页或低门槛评估入口。
  3. 弱意向：用户只是围观、质疑、玩梗或信息不足，先用利益钩子和低门槛问题引导回复。
- 暂时不要发送优惠券或企业名片；需要承接时优先使用资料、评估、预约、草料码或直接联系方式。
- 如果页面目标指定微信号、微信群、企业微信、vx、wx 等直接联系方式，不要用名片/卡片代替；
  联系方式应直接、自然地写在私信文本里。
- 不要在同一条回复里同时说“回复关键词/回我一句后再发”和“已经发了入口/联系方式”。
- 回复要像真人私信，不要像客服模板或知识库问答。
- 视频概述只用于判断用户可能感兴趣的方向，不要在私信里明说“视频里讲的是/视频介绍的是/看到这个视频”；可以自然表达为“看到您对xx比较感兴趣”，信息不足时也可以不提视频。
- 优先用本地部署大模型；如果页面配置为第三方 API，则按页面配置执行。
- 短对话优先使用完整原始上下文；长对话使用压缩核心记忆，信息不足时再读取原始上下文。
- 输出要围绕项目资料和场景提示词生成，避免跨项目串话。
""".strip()


DEFAULT_PROJECT_NAME = "大健康AI跨境综合企业服务平台"
DEFAULT_KNOWLEDGE_BASE_NAME = "公司私信知识库"
RECRUITMENT_KNOWLEDGE_BASE_NAME = "招聘知识库"


DEFAULT_KNOWLEDGE_DOCUMENTS: List[Dict[str, Any]] = [
    {
        "doc_id": "company_positioning",
        "title": "公司定位",
        "knowledge_base": DEFAULT_KNOWLEDGE_BASE_NAME,
        "domain": "公司定位",
        "section": "默认板块",
        "category": "公司定位",
        "sender_identity": "品牌客服",
        "relative_path": "knowledge/files/公司私信知识库/公司定位/默认板块/公司定位.md",
        "description": "说明公司总体定位、服务对象、业务闭环和五大聚焦方向。",
        "keywords": ["公司", "定位", "介绍", "平台", "大健康", "AI", "跨境", "服务对象", "解决方案"],
        "content": """# 公司定位

公司是一家以大健康为核心、AI技术为驱动、跨境业务为延伸的综合型企业服务平台，聚焦关节健康、干细胞合规医疗、AI内容生产、智能硬件、跨境电商五大方向，打造“技术+产品+服务+渠道”全链路商业闭环，面向C端消费者、B端商户、G端公立机构提供产品与解决方案。
""".strip(),
    },
    {
        "doc_id": "healthcare_medical",
        "title": "大健康医疗板块",
        "knowledge_base": DEFAULT_KNOWLEDGE_BASE_NAME,
        "domain": "三大核心业务板块",
        "section": "默认板块",
        "category": "三大核心业务板块",
        "sender_identity": "健康顾问助理",
        "relative_path": "knowledge/files/公司私信知识库/三大核心业务板块/默认板块/大健康医疗板块.md",
        "description": "大健康医疗核心主业，覆盖关节健康、干细胞、PRP、产品体系、合作渠道与合规优势。",
        "keywords": [
            "大健康", "医疗", "关节", "膝骨关节", "骨积液", "软骨修复", "口服", "贴剂",
            "注射", "理疗仪", "非变二型骨胶原肽", "干细胞", "PRP", "医院", "门诊",
            "诊所", "康养", "养老", "社区健康",
        ],
        "content": """# 大健康医疗板块（核心主业）

- 聚焦膝骨关节、骨积液、软骨修复等垂直领域。
- 产品体系：口服、贴剂、注射、理疗仪、非变二型骨胶原肽等。
- 高端项目：合规干细胞、PRP注射、临床治疗服务。
- 合作渠道：医院、门诊、诊所、康养机构、公办养老、社区健康项目。
- 合规优势：资质齐全、自有门诊、临床落地、国资背景资源支持。
""".strip(),
    },
    {
        "doc_id": "ai_technology_hardware",
        "title": "AI技术与智能硬件板块",
        "knowledge_base": DEFAULT_KNOWLEDGE_BASE_NAME,
        "domain": "三大核心业务板块",
        "section": "默认板块",
        "category": "三大核心业务板块",
        "sender_identity": "运营顾问",
        "relative_path": "knowledge/files/公司私信知识库/三大核心业务板块/默认板块/AI技术与智能硬件板块.md",
        "description": "AI内容生产、智能硬件、私信引流、舆论管控、数据采集和AI Agent自动化调度能力。",
        "keywords": [
            "AI", "智能硬件", "内容生产", "文生视频", "批量剪辑", "自动发布", "评论采集",
            "数据统计", "录音背夹", "可穿戴", "录音转纪要", "远程小主机", "边缘计算",
            "私信", "引私域", "舆论管控", "分布式爬虫", "千万级数据", "Agent", "自动化调度",
        ],
        "content": """# AI技术与智能硬件板块

- AI内容生产：文生视频、批量剪辑、自动发布、评论采集、数据统计。
- 智能产品：录音背夹/智能可穿戴设备录音转纪要建议协助办公等，远程小主机（边缘计算），本地+云端协同。
- 核心系统：数据收集-自动私信引私域系统、舆论管控系统等。
- 技术能力：分布式爬虫、千万级数据优化、AI Agent自动化调度。
""".strip(),
    },
    {
        "doc_id": "cross_border_services",
        "title": "跨境与综合业务板块",
        "knowledge_base": DEFAULT_KNOWLEDGE_BASE_NAME,
        "domain": "三大核心业务板块",
        "section": "默认板块",
        "category": "三大核心业务板块",
        "sender_identity": "跨境运营顾问",
        "relative_path": "knowledge/files/公司私信知识库/三大核心业务板块/默认板块/跨境与综合业务板块.md",
        "description": "跨境电商、AI选品、内容托管、自动化运营、留学移民和企业资源对接。",
        "keywords": [
            "跨境", "综合业务", "跨境电商", "全平台运营", "AI选品", "内容托管",
            "自动化运营", "留学", "移民", "企业服务", "资源对接",
        ],
        "content": """# 跨境与综合业务板块

- 跨境电商：全平台运营、AI选品、内容托管、自动化运营。
- 跨境服务：留学移民项目批发、企业服务、资源对接。
""".strip(),
    },
    {
        "doc_id": "recruitment_company_profile",
        "title": "招聘公司介绍",
        "knowledge_base": RECRUITMENT_KNOWLEDGE_BASE_NAME,
        "domain": "公司资料",
        "section": "默认板块",
        "category": "公司资料",
        "sender_identity": "招聘助理",
        "relative_path": "knowledge/files/招聘知识库/公司资料/默认板块/招聘公司介绍.md",
        "description": "当候选人询问公司背景、业务方向、团队情况或岗位可信度时读取。",
        "keywords": ["招聘", "公司", "业务", "团队", "岗位", "候选人", "HR", "介绍", "发展方向"],
        "content": """# 招聘公司介绍

公司是一家以大健康、AI技术、智能硬件和跨境业务为核心方向的综合型企业服务平台。招聘沟通中可以强调公司正在搭建“内容获客、自动化运营、客户承接、业务转化”的完整链路，岗位会接触真实业务项目、AI工具和跨部门协作。

对候选人表达时，应保持真实克制：可以说明业务方向、团队需要、成长空间和面试安排，不夸大规模、融资、薪资或确定录用结果。
""".strip(),
    },
    {
        "doc_id": "recruitment_ai_project_manager",
        "title": "AI项目经理岗位说明",
        "knowledge_base": RECRUITMENT_KNOWLEDGE_BASE_NAME,
        "domain": "岗位资料",
        "section": "AI项目经理",
        "category": "岗位资料",
        "sender_identity": "招聘助理",
        "relative_path": "knowledge/files/招聘知识库/岗位资料/AI项目经理/AI项目经理岗位说明.md",
        "description": "当候选人评论或咨询 AI 项目经理、产品项目、自动化系统落地岗位时读取。",
        "keywords": ["招聘", "AI项目经理", "项目经理", "产品", "需求", "自动化", "沟通", "交付", "岗位职责"],
        "content": """# AI项目经理岗位说明

岗位方向：负责 AI 私信承接、评论采集、内容生产、数据看板、文件解析入库等项目的需求梳理和落地协调。

主要职责：
- 和业务方沟通需求，整理流程、页面、接口和提示词规则。
- 协调开发、运营、测试推进项目上线。
- 跟进用户反馈，持续优化智能体回复效果和后台管理体验。

适合人群：有产品/项目经验、理解 AI 工具、有较强沟通推进能力，能把模糊业务需求拆成可执行任务。
""".strip(),
    },
    {
        "doc_id": "recruitment_engineer_role",
        "title": "技术工程师岗位说明",
        "knowledge_base": RECRUITMENT_KNOWLEDGE_BASE_NAME,
        "domain": "岗位资料",
        "section": "技术工程师",
        "category": "岗位资料",
        "sender_identity": "招聘助理",
        "relative_path": "knowledge/files/招聘知识库/岗位资料/技术工程师/技术工程师岗位说明.md",
        "description": "当候选人咨询工程师、后端、前端、自动化、爬虫或 AI Agent 开发岗位时读取。",
        "keywords": ["招聘", "工程师", "后端", "前端", "自动化", "爬虫", "AI Agent", "接口", "数据库", "岗位"],
        "content": """# 技术工程师岗位说明

岗位方向：参与本地智能体系统、后台管理页面、第三方模型 API、文件解析入库、招聘自动化和评论私信自动回复等模块开发。

主要职责：
- 开发和维护 Python 后端接口、Vue 管理页和本地数据存储。
- 接入 Minimax、DeepSeek、Ollama 等模型供应商。
- 优化知识库检索、prompt 组装、会话上下文和文件解析流程。

适合人群：熟悉 Python 或前端工程，愿意快速理解业务流程，能独立排查接口、浏览器自动化或数据结构问题。
""".strip(),
    },
    {
        "doc_id": "recruitment_interview_process",
        "title": "面试流程说明",
        "knowledge_base": RECRUITMENT_KNOWLEDGE_BASE_NAME,
        "domain": "招聘流程",
        "section": "默认板块",
        "category": "招聘流程",
        "sender_identity": "招聘助理",
        "relative_path": "knowledge/files/招聘知识库/招聘流程/默认板块/面试流程说明.md",
        "description": "当候选人询问如何投递、是否方便沟通、面试安排、到岗时间时读取。",
        "keywords": ["招聘", "面试", "投递", "简历", "沟通", "到岗", "流程", "安排", "HR"],
        "content": """# 面试流程说明

默认流程：
1. 先确认候选人的岗位方向、过往经验、期望城市/薪资和到岗时间。
2. 引导候选人发送简历或作品信息。
3. 初步匹配后安排线上或线下面试。
4. 面试后根据业务负责人反馈推进复试、试岗或 offer 沟通。

私信口吻：不要一上来压迫候选人投简历，先说明看到对方对岗位有兴趣，再问“方便简单了解下你目前主要做哪块吗？”。
""".strip(),
    },
    {
        "doc_id": "recruitment_compensation_compliance",
        "title": "薪资福利与沟通合规",
        "knowledge_base": RECRUITMENT_KNOWLEDGE_BASE_NAME,
        "domain": "薪资福利与合规",
        "section": "默认板块",
        "category": "薪资福利与合规",
        "sender_identity": "招聘助理",
        "relative_path": "knowledge/files/招聘知识库/薪资福利与合规/默认板块/薪资福利与沟通合规.md",
        "description": "当候选人询问薪资、福利、真假、录用承诺或敏感招聘问题时读取。",
        "keywords": ["招聘", "薪资", "福利", "待遇", "社保", "真假", "录用", "合规", "承诺", "面试"],
        "content": """# 薪资福利与沟通合规

可表达内容：
- 薪资以岗位预算和面试结果为准，具体范围可在确认岗位后沟通。
- 福利、工作时间、社保、试用期等以正式招聘说明和公司制度为准。
- 可以邀请候选人先发简历或补充经验，再判断是否匹配。

禁止表达：
- 不承诺“肯定录用”“保底高薪”“无需面试直接入职”。
- 不虚构公司规模、融资、背书、岗位数量。
- 不索要身份证、银行卡等敏感隐私信息。
""".strip(),
    },
]


DEFAULT_SCENES: Dict[str, Dict[str, str]] = {
    "auto": {
        "name": "自动匹配场景",
        "prompt": "根据用户评论自动匹配最相关的项目场景资料。",
        "materials": "",
    },
    "health_presales": {
        "name": "大健康医疗售前",
        "prompt": """
场景目标：识别关节健康、膝骨关节、骨积液、软骨修复、康养、干细胞、PRP 等意向客户。
回复策略：先共情疼痛/行动不便/恢复焦虑，再给一个低门槛体验或资料钩子。
意向分档：问效果/适用/怎么买/多少钱/想试试为强意向，给【专属评估入口】或【预约咨询入口】；只说疼痛或担心为中意向，先给【草料码领取页】或【关节养护资料页】；泛泛围观先引导回复。
优先物料：使用【专属评估入口】、【预约咨询入口】、【草料码领取页】、【关节养护资料页】。
合规要求：不承诺疗效，不虚构医生、医院、名额、价格、国资背书；涉及治疗用“需专业评估后确认”。
""".strip(),
        "materials": """
公司定位：
公司是一家以大健康为核心、AI技术为驱动、跨境业务为延伸的综合型企业服务平台，打造“技术+产品+服务+渠道”全链路商业闭环。

大健康医疗板块：
- 聚焦膝骨关节、骨积液、软骨修复等垂直领域。
- 产品体系：口服、贴剂、注射、理疗仪、非变二型骨胶原肽等。
- 高端项目：合规干细胞、PRP注射、临床治疗服务。
- 合作渠道：医院、门诊、诊所、康养机构、公办养老、社区健康项目。
- 合规优势：资质齐全、自有门诊、临床落地、国资背景资源支持。
""".strip(),
    },
    "ai_hardware_presales": {
        "name": "AI技术与智能硬件售前",
        "prompt": """
场景目标：识别内容生产、自动发布、评论采集、数据统计、智能硬件、录音纪要、边缘计算、私信引流等需求。
回复策略：强调节省人工、自动化闭环、先给演示/体验入口。
优先物料：使用【平台AI工具体验入口】、【草料码演示页】、【方案资料页】。
""".strip(),
        "materials": """
AI技术与智能硬件板块：
- AI内容生产：文生视频、批量剪辑、自动发布、评论采集、数据统计。
- 智能产品：录音背夹、智能可穿戴设备、录音转纪要建议协助办公、远程小主机、本地+云端协同。
- 核心系统：数据收集、自动私信引私域系统、舆论管控系统。
- 技术能力：分布式爬虫、千万级数据优化、AI Agent自动化调度。
""".strip(),
    },
    "cross_border_presales": {
        "name": "跨境综合业务售前",
        "prompt": """
场景目标：识别跨境电商、AI选品、内容托管、自动化运营、留学移民、企业资源对接等需求。
回复策略：强调方案对接、资源入口、先发资料卡片让用户继续沟通。
优先物料：使用【公司方案资料卡】、【专属预约入口】、【资源对接入口】。
""".strip(),
        "materials": """
跨境与综合业务板块：
- 跨境电商：全平台运营、AI选品、内容托管、自动化运营。
- 跨境服务：留学移民项目批发、企业服务、资源对接。
""".strip(),
    },
    "general_presales": {
        "name": "综合企业服务售前",
        "prompt": """
场景目标：当用户意图不明确时，先判断其更接近大健康、AI技术、智能硬件或跨境服务。
回复策略：用轻量问题或综合入口承接，不要强行推具体疗效、价格、名额。
意向分档：强意向给预约入口、评估入口或资料页；中意向给资料页/入口；弱意向先问一个低门槛问题。
优先物料：使用【专属体验入口】、【资料领取页】、【预约咨询入口】。
""".strip(),
        "materials": """
综合定位：
面向C端消费者、B端商户、G端公立机构提供产品与解决方案，覆盖大健康医疗、AI技术与智能硬件、跨境与综合业务。
""".strip(),
    },
}


@dataclass
class ProjectBundle:
    project_id: str
    project_name: str
    scene_id: str
    scene_name: str
    global_prompt: str
    scene_prompt: str
    materials_context: str
    document_context: str = ""

    def to_prompt_context(self) -> str:
        return f"""
<global_operator_prompt>
{self.global_prompt}
</global_operator_prompt>

<project>
project_id: {self.project_id}
project_name: {self.project_name}
scene_id: {self.scene_id}
scene_name: {self.scene_name}
</project>

<project_scene_prompt>
{self.scene_prompt}
</project_scene_prompt>

<project_materials>
{self.materials_context}
</project_materials>

<selected_project_documents>
{self.document_context or "No project documents were selected."}
</selected_project_documents>
""".strip()


class ProjectMaterialStore:
    """File-backed project/scene prompt and material store for the local web agent."""

    def __init__(self, base_dir: Optional[Path] = None):
        root = Path(__file__).resolve().parents[2]
        self.base_dir = base_dir or root / "content" / "session_rag_projects"
        self.base_dir.mkdir(parents=True, exist_ok=True)
        self.default_project_id = self.ensure_default_project()
        for project_dir in self._project_dirs():
            self.ensure_default_knowledge_documents(project_dir, create_missing_defaults=False)

    @staticmethod
    def _knowledge_base_id(name: str) -> str:
        return "kb_" + uuid.uuid5(uuid.NAMESPACE_URL, str(name or DEFAULT_KNOWLEDGE_BASE_NAME)).hex[:10]

    def ensure_default_project(self) -> str:
        existing = self._project_dirs()
        if existing:
            return existing[0].name

        project_id = self._generated_project_id(DEFAULT_PROJECT_NAME)
        project_dir = self._write_project_template(project_id, DEFAULT_PROJECT_NAME)
        self.ensure_default_knowledge_documents(project_dir, create_missing_defaults=True)
        return project_id

    def create_project(self, name: str) -> Dict[str, Any]:
        clean_name = str(name or "").strip() or "新项目"
        project_id = self._generated_project_id(clean_name)
        if (self.base_dir / project_id).exists():
            project_id = f"{project_id}_{uuid.uuid4().hex[:6]}"
        project_dir = self._write_project_template(project_id, clean_name)
        self.ensure_default_knowledge_documents(project_dir, create_missing_defaults=True)
        return {
            "project_id": project_id,
            "name": clean_name,
            "default_scene": "auto",
            "path": str(project_dir),
        }

    def _write_project_template(self, project_id: str, name: str) -> Path:
        project_dir = self.base_dir / project_id
        scenes_dir = project_dir / "scenes"
        scenes_dir.mkdir(parents=True, exist_ok=True)
        (project_dir / "project.json").write_text(
            json.dumps(
                {
                    "project_id": project_id,
                    "name": name,
                    "default_scene": "auto",
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        (project_dir / "global_prompt.md").write_text(GLOBAL_OPERATOR_PROMPT + "\n", encoding="utf-8")

        for scene_id, scene in DEFAULT_SCENES.items():
            scene_dir = scenes_dir / scene_id
            scene_dir.mkdir(parents=True, exist_ok=True)
            (scene_dir / "scene.json").write_text(
                json.dumps({"scene_id": scene_id, "name": scene["name"]}, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            (scene_dir / "prompt.md").write_text(scene["prompt"] + "\n", encoding="utf-8")
            (scene_dir / "materials.md").write_text(scene["materials"] + "\n", encoding="utf-8")

        return project_dir

    def list_projects(self) -> List[Dict[str, Any]]:
        projects = []
        for project_dir in self._project_dirs():
            meta = self._read_json(project_dir / "project.json")
            project_id = project_dir.name
            projects.append(
                {
                    "project_id": project_id,
                    "name": meta.get("name") or project_id,
                    "default_scene": meta.get("default_scene") or "auto",
                    "scenes": self._list_scenes(project_dir),
                }
            )
        return projects

    @staticmethod
    def _iter_description_documents(descriptions: Dict[str, Any]):
        if not isinstance(descriptions, dict):
            return
        knowledge_bases = descriptions.get("knowledge_bases")
        if not isinstance(knowledge_bases, list) or not knowledge_bases:
            knowledge_bases = [
                {
                    "name": DEFAULT_KNOWLEDGE_BASE_NAME,
                    "domains": descriptions.get("domains") if isinstance(descriptions.get("domains"), list) else [],
                }
            ]
        for base in knowledge_bases:
            if not isinstance(base, dict):
                continue
            knowledge_base_name = str(base.get("name") or DEFAULT_KNOWLEDGE_BASE_NAME)
            for domain in base.get("domains") or []:
                if not isinstance(domain, dict):
                    continue
                domain_name = str(domain.get("name") or "默认领域")
                for section in domain.get("sections") or []:
                    if not isinstance(section, dict):
                        continue
                    section_name = str(section.get("name") or "默认板块")
                    for item in section.get("documents") or []:
                        if isinstance(item, dict):
                            yield knowledge_base_name, domain_name, section_name, item

    def ensure_default_knowledge_documents(self, project_dir: Path, create_missing_defaults: bool = False) -> None:
        knowledge_dir = project_dir / "knowledge"
        knowledge_dir.mkdir(parents=True, exist_ok=True)
        existing_manifest = self._read_json(knowledge_dir / "manifest.json")
        deleted_doc_ids = set(existing_manifest.get("deleted_doc_ids") or [])
        existing_docs = {
            item.get("doc_id"): item
            for item in existing_manifest.get("documents", [])
            if item.get("doc_id")
        }

        documents = []
        for doc in DEFAULT_KNOWLEDGE_DOCUMENTS:
            if doc["doc_id"] in deleted_doc_ids:
                continue
            merged = {
                **{key: value for key, value in doc.items() if key != "content"},
                **existing_docs.get(doc["doc_id"], {}),
            }
            _, resolved_relative = self._read_knowledge_document_text(project_dir, merged)
            doc_exists = bool(resolved_relative and (project_dir / resolved_relative).is_file())
            if not doc_exists:
                if not create_missing_defaults:
                    if doc["doc_id"] in existing_docs:
                        deleted_doc_ids.add(doc["doc_id"])
                    continue
                doc_path = project_dir / doc["relative_path"]
                doc_path.parent.mkdir(parents=True, exist_ok=True)
                doc_path.write_text(doc["content"] + "\n", encoding="utf-8")
                resolved_relative = doc["relative_path"]

            merged["knowledge_base"] = doc.get("knowledge_base") or merged.get("knowledge_base") or DEFAULT_KNOWLEDGE_BASE_NAME
            merged["kb_id"] = merged.get("kb_id") or self._knowledge_base_id(merged["knowledge_base"])
            merged["domain"] = doc.get("domain") or merged.get("domain") or merged.get("category") or "默认领域"
            merged["section"] = doc.get("section") or merged.get("section") or "默认板块"
            merged["category"] = doc.get("category") or merged.get("category") or merged["domain"]
            merged["relative_path"] = resolved_relative or merged.get("relative_path") or doc.get("relative_path")
            if not self._is_valid_sender_identity(merged.get("sender_identity")):
                merged["sender_identity"] = doc.get("sender_identity") or self._suggest_sender_identity_by_rules(
                    merged.get("category", ""),
                    merged.get("title", ""),
                    " ".join(str(keyword) for keyword in merged.get("keywords", [])),
                )
            documents.append(merged)

        custom_docs = [
            item
            for item in existing_manifest.get("documents", [])
            if item.get("doc_id") not in {doc["doc_id"] for doc in DEFAULT_KNOWLEDGE_DOCUMENTS}
            and item.get("doc_id") not in deleted_doc_ids
            and self._manifest_document_file_exists(project_dir, item)
        ]
        for item in custom_docs:
            knowledge_base = str(item.get("knowledge_base") or DEFAULT_KNOWLEDGE_BASE_NAME).strip() or DEFAULT_KNOWLEDGE_BASE_NAME
            item["knowledge_base"] = knowledge_base
            item["kb_id"] = item.get("kb_id") or self._knowledge_base_id(knowledge_base)
        documents.extend(custom_docs)

        seen_doc_ids = {item.get("doc_id") for item in documents if item.get("doc_id")}
        descriptions = self._read_json(knowledge_dir / "document_descriptions.json")
        for knowledge_base_name, domain_name, section_name, item in self._iter_description_documents(descriptions):
            doc_id = item.get("doc_id")
            relative_path = item.get("relative_path")
            if not doc_id or doc_id in seen_doc_ids or doc_id in deleted_doc_ids:
                continue
            _, resolved_relative = self._read_knowledge_document_text(
                project_dir,
                {
                    **item,
                    "knowledge_base": knowledge_base_name,
                    "domain": domain_name,
                    "section": section_name,
                },
            )
            if not resolved_relative or not (project_dir / resolved_relative).is_file():
                continue
            category = item.get("category") or domain_name
            documents.append({
                "doc_id": doc_id,
                "title": item.get("title") or item.get("source_file_name") or relative_path,
                "kb_id": self._knowledge_base_id(knowledge_base_name),
                "knowledge_base": knowledge_base_name,
                "domain": domain_name,
                "section": section_name,
                "category": category,
                "sender_identity": item.get("sender_identity") if self._is_valid_sender_identity(item.get("sender_identity")) else self._suggest_sender_identity_by_rules(
                    f"{knowledge_base_name}/{domain_name}/{section_name}",
                    item.get("title") or item.get("source_file_name") or relative_path,
                    " ".join([
                        str(item.get("description") or ""),
                        str(item.get("summary") or ""),
                        " ".join(str(tag) for tag in item.get("tags") or []),
                    ]),
                ),
                "relative_path": resolved_relative,
                "description": item.get("description") or "",
                "keywords": item.get("tags") or item.get("keywords") or [],
                "source_file_name": item.get("source_file_name") or "",
                "imported_at": item.get("updated_at") or "",
                "content_chars": item.get("content_chars") or 0,
            })
            seen_doc_ids.add(doc_id)

        manifest = {"version": 1, "documents": documents}
        if deleted_doc_ids:
            manifest["deleted_doc_ids"] = sorted(deleted_doc_ids)
        (knowledge_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    def list_knowledge_documents(self, project_id: str = "") -> List[Dict[str, Any]]:
        project_dir = self._resolve_project_dir(project_id)
        self.ensure_default_knowledge_documents(project_dir, create_missing_defaults=False)
        manifest = self._read_json(project_dir / "knowledge" / "manifest.json")
        return [
            {
                **doc,
                "sender_identity": self._display_sender_identity(doc.get("sender_identity")) or "品牌客服",
            }
            for doc in manifest.get("documents", [])
            if isinstance(doc, dict)
        ]

    def import_file_to_knowledge(
        self,
        file_name: str,
        file_data: bytes,
        project_id: str = "",
        llm_tools: Any = None,
        model_conf: Optional[Dict[str, Any]] = None,
    ) -> Dict[str, Any]:
        if not file_name or not file_name.strip():
            raise ValueError("file_name is required")
        if not file_data:
            raise ValueError("file_data is empty")

        project_dir = self._resolve_project_dir(project_id)
        self.ensure_default_knowledge_documents(project_dir, create_missing_defaults=False)

        content = self.parse_file_content(file_name, file_data).strip()
        if not content:
            raise ValueError("file content is empty after parsing")

        meta = self._suggest_imported_document_meta(file_name, content, llm_tools, model_conf)
        doc_id = "imported_" + uuid.uuid5(
            uuid.NAMESPACE_URL,
            f"{project_dir.name}:{file_name}:{content[:500]}:{len(content)}",
        ).hex[:12]
        title = self._safe_title(meta.get("title") or Path(file_name).stem or "导入资料")
        category = str(meta.get("category") or "导入资料").strip() or "导入资料"
        knowledge_base = str(meta.get("knowledge_base") or DEFAULT_KNOWLEDGE_BASE_NAME).strip() or DEFAULT_KNOWLEDGE_BASE_NAME
        domain = str(meta.get("domain") or category).strip() or "导入资料"
        section = str(meta.get("section") or "默认板块").strip() or "默认板块"
        keywords = self._normalize_keywords(meta.get("keywords") or [])
        description = str(meta.get("description") or f"从 {file_name} 解析入库的项目资料").strip()
        sender_identity = self._display_sender_identity(
            meta.get("sender_identity") if self._is_valid_sender_identity(meta.get("sender_identity")) else self._suggest_sender_identity_by_rules(category, file_name, content)
        ) or "品牌客服"

        relative_path = f"knowledge/files/{self._safe_title(knowledge_base)}/{self._safe_title(domain)}/{self._safe_title(section)}/{title}.md"
        doc_path = project_dir / relative_path
        doc_path.parent.mkdir(parents=True, exist_ok=True)
        doc_path.write_text(
            self._format_imported_markdown(
                title=title,
                source_file_name=file_name,
                category=category,
                description=description,
                sender_identity=sender_identity,
                content=content,
            ),
            encoding="utf-8",
        )

        chunks = self._split_text(content)
        self._upsert_manifest_document(
            project_dir,
            {
                "doc_id": doc_id,
                "title": title,
                "category": category,
                "kb_id": self._knowledge_base_id(knowledge_base),
                "knowledge_base": knowledge_base,
                "domain": domain,
                "section": section,
                "sender_identity": sender_identity,
                "relative_path": relative_path,
                "description": description,
                "keywords": keywords,
                "source_file_name": file_name,
                "imported_at": datetime.now().isoformat(timespec="seconds"),
                "chunk_count": len(chunks),
            },
        )
        self._upsert_chunks(project_dir, doc_id, file_name, relative_path, chunks, sender_identity=sender_identity)

        return {
            "doc_id": doc_id,
            "title": title,
            "category": category,
            "kb_id": self._knowledge_base_id(knowledge_base),
            "knowledge_base": knowledge_base,
            "domain": domain,
            "section": section,
            "sender_identity": sender_identity,
            "relative_path": relative_path,
            "absolute_path": str(doc_path),
            "description": description,
            "keywords": keywords,
            "chunk_count": len(chunks),
            "content_chars": len(content),
        }

    def parse_file_content(self, file_name: str, file_data: bytes) -> str:
        ext = Path(file_name).suffix.lower()
        if ext == ".xlsx":
            return self._parse_xlsx_with_stdlib(file_data)
        if ext == ".xls":
            return self._parse_xls_with_optional_libs(file_data)
        if ext == ".pdf":
            return self._parse_pdf_with_optional_libs(file_data)

        try:
            from aisec_agent.logic._tools import FileReaderTools

            content = FileReaderTools().get_file_content(BytesIO(file_data), file_name).strip()
            if content and not self._is_file_reader_error(content):
                return content
        except Exception:
            pass
        return self._fallback_parse_file_content(file_name, file_data)

    @staticmethod
    def _is_file_reader_error(content: str) -> bool:
        error_prefixes = (
            "读取文件时发生错误",
            "无法读取",
            "文件不是有效",
            "Error loading",
            "Error extracting",
        )
        return any(content.startswith(prefix) for prefix in error_prefixes)

    @staticmethod
    def _fallback_parse_file_content(file_name: str, file_data: bytes) -> str:
        ext = Path(file_name).suffix.lower()
        if ext in {".txt", ".md", ".csv", ".json", ".log", ".xml", ".py", ".java", ".js", ".html", ".css"}:
            return ProjectMaterialStore._decode_text_file(file_data)
        if ext == ".docx":
            return ProjectMaterialStore._parse_docx_with_stdlib(file_data)
        if ext == ".xlsx":
            return ProjectMaterialStore._parse_xlsx_with_stdlib(file_data)
        if ext == ".pdf":
            return ProjectMaterialStore._parse_pdf_with_optional_libs(file_data)
        if ext == ".xls":
            return ProjectMaterialStore._parse_xls_with_optional_libs(file_data)
        raise ValueError(f"当前环境无法解析该文件类型: {ext or file_name}")

    @staticmethod
    def _decode_text_file(file_data: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5"):
            try:
                return file_data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_data.decode("utf-8", errors="replace")

    @staticmethod
    def _parse_docx_with_stdlib(file_data: bytes) -> str:
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        xml_names = ["word/document.xml"]
        try:
            with ZipFile(BytesIO(file_data)) as archive:
                names = archive.namelist()
                xml_names.extend(
                    name for name in names
                    if re.match(r"word/(header|footer)\d+\.xml$", name)
                )
                paragraphs: List[str] = []
                for xml_name in dict.fromkeys(xml_names):
                    if xml_name not in names:
                        continue
                    root = ET.fromstring(archive.read(xml_name))
                    for paragraph in root.iter(f"{namespace}p"):
                        parts = []
                        for node in paragraph.iter():
                            if node.tag == f"{namespace}t":
                                parts.append(node.text or "")
                            elif node.tag == f"{namespace}tab":
                                parts.append("\t")
                            elif node.tag in {f"{namespace}br", f"{namespace}cr"}:
                                parts.append("\n")
                        text = "".join(parts).strip()
                        if text:
                            paragraphs.append(text)
                return "\n".join(paragraphs)
        except BadZipFile as e:
            raise ValueError("Word 文件不是有效的 docx 格式") from e
        except ET.ParseError as e:
            raise ValueError("Word 文件 XML 解析失败") from e

    @staticmethod
    def _parse_xlsx_with_stdlib(file_data: bytes) -> str:
        try:
            return ProjectMaterialStore._parse_xlsx_with_openpyxl(file_data)
        except Exception:
            pass

        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        try:
            with ZipFile(BytesIO(file_data)) as archive:
                names = archive.namelist()
                shared_strings = ProjectMaterialStore._read_xlsx_shared_strings(archive, namespace)
                sheet_names = ProjectMaterialStore._read_xlsx_sheet_names(archive)
                output = []
                sheet_files = sorted(
                    name for name in names
                    if re.match(r"xl/worksheets/sheet\d+\.xml$", name)
                )
                for index, sheet_file in enumerate(sheet_files, 1):
                    sheet_name = sheet_names.get(sheet_file, f"Sheet{index}")
                    rows = ProjectMaterialStore._read_xlsx_sheet_rows(archive, sheet_file, namespace, shared_strings)
                    if not rows:
                        continue
                    output.append(ProjectMaterialStore._format_sheet_rows_as_records(sheet_name, rows))
                return "\n".join(output)
        except BadZipFile as e:
            raise ValueError("Excel 文件不是有效的 xlsx 格式") from e
        except ET.ParseError as e:
            raise ValueError("Excel 文件 XML 解析失败") from e

    @staticmethod
    def _parse_xlsx_with_openpyxl(file_data: bytes) -> str:
        from openpyxl import load_workbook  # type: ignore

        workbook = load_workbook(BytesIO(file_data), data_only=True, read_only=False)
        output = []
        for worksheet in workbook.worksheets:
            grid = ProjectMaterialStore._worksheet_values_with_merged_cells(worksheet)
            records = ProjectMaterialStore._format_sheet_rows_as_records(worksheet.title, grid)
            if records:
                output.append(records)
        return "\n\n".join(output)

    @staticmethod
    def _worksheet_values_with_merged_cells(worksheet: Any) -> List[List[str]]:
        rows = []
        for row in worksheet.iter_rows():
            rows.append([ProjectMaterialStore._normalize_cell_value(cell.value) for cell in row])

        for merged_range in worksheet.merged_cells.ranges:
            value = ProjectMaterialStore._normalize_cell_value(
                worksheet.cell(merged_range.min_row, merged_range.min_col).value
            )
            if not value:
                continue
            for row_index in range(merged_range.min_row - 1, merged_range.max_row):
                if row_index >= len(rows):
                    continue
                while len(rows[row_index]) < merged_range.max_col:
                    rows[row_index].append("")
                for col_index in range(merged_range.min_col - 1, merged_range.max_col):
                    rows[row_index][col_index] = value

        return rows

    @staticmethod
    def _format_sheet_rows_as_records(sheet_name: str, rows: List[List[str]]) -> str:
        useful_rows = [row for row in rows if any(cell.strip() for cell in row)]
        if not useful_rows:
            return ""

        header_index = ProjectMaterialStore._detect_header_row(useful_rows)
        headers = ProjectMaterialStore._build_table_headers(useful_rows[header_index])
        output = [f"## {sheet_name}", "表格解析方式：按数据行提取，每一行都是一条记录，字段名来自表头。"]

        for row_number, row in enumerate(useful_rows[header_index + 1:], start=header_index + 2):
            pairs = ProjectMaterialStore._row_to_header_pairs(headers, row)
            if not pairs:
                continue
            output.append(f"\n### {sheet_name} 第 {row_number} 行")
            for header, value in pairs:
                output.append(f"- {header}: {value}")
        return "\n".join(output)

    @staticmethod
    def _detect_header_row(rows: List[List[str]]) -> int:
        best_index = 0
        best_score = -1
        for index, row in enumerate(rows[:10]):
            non_empty = [cell for cell in row if cell.strip()]
            unique = len(set(non_empty))
            score = len(non_empty) + unique
            if score > best_score:
                best_index = index
                best_score = score
        return best_index

    @staticmethod
    def _build_table_headers(row: List[str]) -> List[str]:
        headers = []
        seen: Dict[str, int] = {}
        for index, value in enumerate(row, start=1):
            header = value.strip() or f"列{index}"
            count = seen.get(header, 0) + 1
            seen[header] = count
            if count > 1:
                header = f"{header}_{count}"
            headers.append(header)
        return headers

    @staticmethod
    def _row_to_header_pairs(headers: List[str], row: List[str]) -> List[tuple[str, str]]:
        width = max(len(headers), len(row))
        pairs = []
        for index in range(width):
            header = headers[index] if index < len(headers) else f"列{index + 1}"
            value = row[index] if index < len(row) else ""
            value = value.strip()
            if value:
                pairs.append((header, value))
        return pairs

    @staticmethod
    def _normalize_cell_value(value: Any) -> str:
        if value is None:
            return ""
        text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
        return re.sub(r"\n{3,}", "\n\n", text)

    @staticmethod
    def _read_xlsx_shared_strings(archive: ZipFile, namespace: str) -> List[str]:
        if "xl/sharedStrings.xml" not in archive.namelist():
            return []
        root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
        strings = []
        for item in root.iter(f"{namespace}si"):
            parts = [node.text or "" for node in item.iter(f"{namespace}t")]
            strings.append("".join(parts))
        return strings

    @staticmethod
    def _read_xlsx_sheet_names(archive: ZipFile) -> Dict[str, str]:
        if "xl/workbook.xml" not in archive.namelist() or "xl/_rels/workbook.xml.rels" not in archive.namelist():
            return {}

        workbook_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
        rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_targets = {}
        for rel in rels_root.iter(f"{rel_ns}Relationship"):
            rel_id = rel.attrib.get("Id", "")
            target = rel.attrib.get("Target", "")
            if rel_id and target:
                rel_targets[rel_id] = "xl/" + target.lstrip("/")

        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        sheet_names = {}
        for sheet in workbook_root.iter(f"{workbook_ns}sheet"):
            rel_id = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "")
            target = rel_targets.get(rel_id)
            if target:
                sheet_names[target] = sheet.attrib.get("name", "") or target
        return sheet_names

    @staticmethod
    def _read_xlsx_sheet_rows(
        archive: ZipFile,
        sheet_file: str,
        namespace: str,
        shared_strings: List[str],
    ) -> List[List[str]]:
        root = ET.fromstring(archive.read(sheet_file))
        rows = []
        for row in root.iter(f"{namespace}row"):
            cells = []
            for cell in row.iter(f"{namespace}c"):
                cells.append(ProjectMaterialStore._read_xlsx_cell(cell, namespace, shared_strings))
            while cells and not cells[-1]:
                cells.pop()
            if any(cell.strip() for cell in cells):
                rows.append(cells)
        return rows

    @staticmethod
    def _read_xlsx_cell(cell: ET.Element, namespace: str, shared_strings: List[str]) -> str:
        cell_type = cell.attrib.get("t", "")
        if cell_type == "inlineStr":
            return "".join(node.text or "" for node in cell.iter(f"{namespace}t")).strip()

        value = cell.find(f"{namespace}v")
        raw = value.text if value is not None else ""
        if cell_type == "s":
            try:
                return shared_strings[int(raw)]
            except (ValueError, IndexError):
                return raw or ""
        if cell_type == "b":
            return "TRUE" if raw == "1" else "FALSE" if raw == "0" else raw or ""
        return raw or ""

    @staticmethod
    def _parse_pdf_with_optional_libs(file_data: bytes) -> str:
        text = ""
        try:
            import fitz  # type: ignore

            with fitz.open(stream=file_data, filetype="pdf") as document:
                text = "\n\n".join(page.get_text() for page in document).strip()
            if text:
                return text
        except Exception:
            pass

        try:
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(BytesIO(file_data))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
            if text:
                return text
        except Exception:
            pass

        try:
            import pdfplumber  # type: ignore

            with pdfplumber.open(BytesIO(file_data)) as pdf:
                text = "\n\n".join(page.extract_text() or "" for page in pdf.pages).strip()
            if text:
                return text
        except Exception:
            pass

        ocr_text = ProjectMaterialStore._ocr_pdf_pages(file_data).strip()
        if ocr_text:
            return ocr_text
        raise ValueError("PDF 没有可提取文本，OCR 也未识别到内容；请确认文件不是图片质量过低或加密保护")

    @staticmethod
    def _ocr_pdf_pages(file_data: bytes, max_pages: int = 12, zoom: float = 2.0) -> str:
        try:
            return ProjectMaterialStore._ocr_pdf_pages_with_paddle(file_data, max_pages=max_pages, zoom=zoom)
        except Exception:
            return ProjectMaterialStore._ocr_pdf_pages_with_rapidocr(file_data, max_pages=max_pages, zoom=zoom)

    @staticmethod
    def _ocr_pdf_pages_with_paddle(file_data: bytes, max_pages: int = 12, zoom: float = 3.0) -> str:
        try:
            import fitz  # type: ignore
            import numpy as np  # type: ignore
            from PIL import Image  # type: ignore
            from paddleocr import PaddleOCR  # type: ignore
        except Exception as e:
            raise ValueError("PDF 没有可提取文本，且当前环境缺少 OCR 依赖 paddleocr") from e

        ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
        page_texts = []
        with fitz.open(stream=file_data, filetype="pdf") as document:
            for page_index, page in enumerate(document):
                if page_index >= max_pages:
                    page_texts.append(f"[OCR 已达到前 {max_pages} 页限制，后续页面未解析]")
                    break
                matrix = fitz.Matrix(zoom, zoom)
                pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                result = ocr.ocr(np.array(image), cls=True)
                lines = ProjectMaterialStore._flatten_paddle_ocr_result(result)
                if lines:
                    page_texts.append(f"--- Page {page_index + 1} OCR ---\n" + "\n".join(lines))
        return "\n\n".join(page_texts)

    @staticmethod
    def _flatten_paddle_ocr_result(result: Any) -> List[str]:
        lines = []
        for page_result in result or []:
            for item in page_result or []:
                if not isinstance(item, (list, tuple)) or len(item) < 2:
                    continue
                text_score = item[1]
                if not isinstance(text_score, (list, tuple)) or not text_score:
                    continue
                text = str(text_score[0]).strip()
                if text:
                    lines.append(text)
        return lines

    @staticmethod
    def _ocr_pdf_pages_with_rapidocr(file_data: bytes, max_pages: int = 12, zoom: float = 3.0) -> str:
        try:
            import fitz  # type: ignore
            import numpy as np  # type: ignore
            from PIL import Image  # type: ignore
            from rapidocr_onnxruntime import RapidOCR  # type: ignore
        except Exception as e:
            raise ValueError("PDF 没有可提取文本，且当前环境缺少 OCR 依赖 rapidocr-onnxruntime") from e

        ocr = RapidOCR()
        page_texts = []
        with fitz.open(stream=file_data, filetype="pdf") as document:
            for page_index, page in enumerate(document):
                if page_index >= max_pages:
                    page_texts.append(f"[OCR 已达到前 {max_pages} 页限制，后续页面未解析]")
                    break
                matrix = fitz.Matrix(zoom, zoom)
                pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                result, _ = ocr(np.array(image))
                lines = []
                for item in result or []:
                    if len(item) >= 2 and item[1]:
                        lines.append(str(item[1]).strip())
                page_text = "\n".join(line for line in lines if line)
                if page_text:
                    page_texts.append(f"--- Page {page_index + 1} OCR ---\n{page_text}")
        return "\n\n".join(page_texts)

    @staticmethod
    def _parse_xls_with_optional_libs(file_data: bytes) -> str:
        try:
            import pandas as pd  # type: ignore

            sheets = pd.read_excel(BytesIO(file_data), sheet_name=None, header=None)
            output = []
            for sheet_name, frame in sheets.items():
                output.append(f"## {sheet_name}")
                for row in frame.fillna("").astype(str).values.tolist():
                    line = "\t".join(cell for cell in row if cell).strip()
                    if line:
                        output.append(line)
            return "\n".join(output)
        except Exception as e:
            raise ValueError("当前环境缺少 xls 解析依赖，请安装 pandas 和 xlrd，或将文件另存为 xlsx 后重试") from e

    def _suggest_imported_document_meta(
        self,
        file_name: str,
        content: str,
        llm_tools: Any = None,
        model_conf: Optional[Dict[str, Any]] = None,
    ) -> Dict[str, Any]:
        fallback = self._suggest_imported_document_meta_by_rules(file_name, content)
        if not llm_tools or not model_conf:
            return fallback
        try:
            prompt = """
你是项目资料入库助手。请根据文件名和正文内容，为这份资料生成入库元数据。

要求：
- title：简短中文标题，不超过30字，不要带文件后缀。
- knowledge_base：从这些知识库中选一个：公司私信知识库、招聘知识库。
- domain：文档所属领域/业务板块，例如 公司定位、三大核心业务板块、大健康医疗、AI技术与智能硬件、跨境与综合业务、公司资料、岗位资料、招聘流程、薪资福利与合规。
- section：文档所属子目录/板块，例如 默认板块、AI项目经理、技术工程师、售前资料、合规边界、产品说明。
- category：从这些类别中选一个：公司定位、三大核心业务板块、大健康医疗、AI技术与智能硬件、跨境与综合业务、优惠活动、话术规则、合规边界、导入资料。
- description：一句话说明这份资料什么时候应该被读取。
- keywords：5到12个关键词，用于后续根据用户评论匹配资料。
- sender_identity：根据资料所属业务建议私信时使用的人员身份，例如 健康顾问助理、运营顾问、跨境运营顾问、招聘助理、品牌客服。

只返回 JSON，字段固定为：title, knowledge_base, domain, section, category, description, keywords, sender_identity。
""".strip()
            message = f"文件名：{file_name}\n\n正文预览：\n{content[:6000]}"
            chat = getattr(llm_tools, model_conf["func_name"])
            response = chat(
                url=model_conf["url"],
                api_key=model_conf.get("key", ""),
                prompt=prompt,
                message=message,
                model=model_conf["model_name"],
                json_format=True,
                stream=False,
                max_len_input=model_conf.get("max_len_input", 16000),
            )
            data = self._parse_json_response(response)
            return {
                "title": data.get("title") or fallback["title"],
                "knowledge_base": data.get("knowledge_base") or fallback["knowledge_base"],
                "domain": data.get("domain") or fallback["domain"],
                "section": data.get("section") or fallback["section"],
                "category": data.get("category") or fallback["category"],
                "description": data.get("description") or fallback["description"],
                "keywords": data.get("keywords") or fallback["keywords"],
                "sender_identity": data.get("sender_identity") or fallback["sender_identity"],
            }
        except Exception:
            return fallback

    def _suggest_imported_document_meta_by_rules(self, file_name: str, content: str) -> Dict[str, Any]:
        text = f"{file_name}\n{content[:8000]}".lower()
        title = Path(file_name).stem or "导入资料"
        knowledge_base = DEFAULT_KNOWLEDGE_BASE_NAME
        category = "导入资料"
        domain = "导入资料"
        section = "默认板块"
        description = f"从 {file_name} 解析入库的项目资料。"

        category_rules = [
            ("岗位资料", ["招聘", "岗位", "候选人", "简历", "面试", "薪资", "boss", "hr", "工程师", "项目经理", "求职", "到岗"]),
            ("大健康医疗", ["关节", "膝", "骨积液", "软骨", "干细胞", "prp", "康养", "理疗", "骨胶原"]),
            ("AI技术与智能硬件", ["ai", "智能硬件", "文生视频", "批量剪辑", "自动发布", "评论采集", "录音", "私信", "舆论", "agent"]),
            ("跨境与综合业务", ["跨境", "电商", "选品", "托管", "留学", "移民", "资源对接"]),
            ("优惠活动", ["优惠", "券", "折扣", "体验", "名额", "领取", "活动"]),
            ("话术规则", ["话术", "首次私信", "私信回复", "钩子", "引导回复"]),
            ("合规边界", ["合规", "禁止", "不能", "不得", "风险", "承诺", "医疗广告"]),
            ("公司定位", ["公司定位", "公司介绍", "综合型企业", "服务平台", "业务板块"]),
        ]
        for candidate, keywords in category_rules:
            if any(keyword in text for keyword in keywords):
                category = candidate
                domain = candidate
                description = f"当用户问题涉及{candidate}相关内容时读取。"
                break
        if category == "岗位资料":
            knowledge_base = RECRUITMENT_KNOWLEDGE_BASE_NAME
            if re.search(r"项目经理|产品|需求|交付", text, flags=re.I):
                section = "AI项目经理"
            elif re.search(r"工程师|后端|前端|开发|爬虫|接口|agent", text, flags=re.I):
                section = "技术工程师"
            elif re.search(r"薪资|待遇|福利|社保|真假|承诺", text, flags=re.I):
                domain = "薪资福利与合规"
            elif re.search(r"面试|简历|投递|到岗|流程", text, flags=re.I):
                domain = "招聘流程"

        keywords = self._extract_keywords_for_manifest(file_name, content)
        return {
            "title": title,
            "knowledge_base": knowledge_base,
            "domain": domain,
            "section": section,
            "category": category,
            "description": description,
            "keywords": keywords,
            "sender_identity": self._suggest_sender_identity_by_rules(category, file_name, content),
        }

    @staticmethod
    def _is_valid_sender_identity(identity: Any) -> bool:
        text = str(identity or "").strip()
        return bool(text) and "?" not in text and "�" not in text

    @staticmethod
    def _display_sender_identity(identity: Any) -> str:
        text = str(identity or "").strip()
        if not text:
            return ""
        cleaned = re.sub(r"(?i)^\s*xx[\s._-]*", "", text).strip()
        return cleaned or text

    @staticmethod
    def _suggest_sender_identity_by_rules(category: str, file_name: str = "", content: str = "") -> str:
        text = f"{category}\n{file_name}\n{content[:8000]}".lower()
        if re.search(r"招聘|岗位|候选人|简历|面试|薪资|boss|hr|求职|到岗", text, flags=re.I):
            return "招聘助理"
        if re.search(r"睡眠|关节|膝|大健康|医疗|骨积液|软骨|干细胞|prp|康养|理疗|健康", text, flags=re.I):
            return "健康顾问助理"
        if re.search(r"文生视频|ai|智能硬件|自动发布|评论采集|私信|内容生产|剪辑|运营|agent|舆论", text, flags=re.I):
            return "运营顾问"
        if re.search(r"跨境|电商|选品|店铺|托管|留学|移民|海外", text, flags=re.I):
            return "跨境运营顾问"
        return "品牌客服"

    @staticmethod
    def _safe_title(title: str) -> str:
        clean = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", str(title)).strip(" ._")
        return clean[:60] or "导入资料"

    @staticmethod
    def _normalize_keywords(keywords: List[Any]) -> List[str]:
        normalized = []
        for keyword in keywords:
            text = str(keyword).strip()
            if text and text not in normalized:
                normalized.append(text)
        return normalized[:12]

    def _extract_keywords_for_manifest(self, file_name: str, content: str) -> List[str]:
        seed_keywords = [
            "公司", "定位", "大健康", "医疗", "关节", "骨积液", "软骨", "干细胞", "PRP",
            "AI", "智能硬件", "内容生产", "自动发布", "评论采集", "私信", "舆论管控",
            "跨境", "电商", "选品", "留学", "移民", "优惠", "体验", "名片", "草料码",
            "话术", "合规", "招聘", "岗位", "候选人", "简历", "面试", "薪资", "福利", "HR",
        ]
        text = f"{file_name}\n{content[:10000]}"
        keywords = [keyword for keyword in seed_keywords if keyword.lower() in text.lower()]
        title_parts = re.split(r"[\s_\-\.]+", Path(file_name).stem)
        keywords.extend(part for part in title_parts if part)
        return self._normalize_keywords(keywords)

    @staticmethod
    def _format_imported_markdown(
        title: str,
        source_file_name: str,
        category: str,
        description: str,
        sender_identity: str,
        content: str,
    ) -> str:
        return f"""# {title}

> 来源文件：{source_file_name}
> 分类：{category}
> 读取说明：{description}
> 建议私信身份：{sender_identity or '品牌客服'}

{content.strip()}
""".strip() + "\n"

    @staticmethod
    def _split_text(text: str, chunk_size: int = 900, overlap: int = 120) -> List[str]:
        clean = re.sub(r"\n{3,}", "\n\n", text.strip())
        if not clean:
            return []
        paragraphs = [part.strip() for part in re.split(r"\n\s*\n", clean) if part.strip()]
        chunks: List[str] = []
        current = ""
        for paragraph in paragraphs:
            if len(paragraph) > chunk_size:
                if current:
                    chunks.append(current.strip())
                    current = ""
                start = 0
                while start < len(paragraph):
                    chunks.append(paragraph[start:start + chunk_size].strip())
                    start += max(1, chunk_size - overlap)
                continue
            if current and len(current) + len(paragraph) + 2 > chunk_size:
                chunks.append(current.strip())
                current = current[-overlap:] + "\n\n" + paragraph if overlap and len(current) > overlap else paragraph
            else:
                current = paragraph if not current else current + "\n\n" + paragraph
        if current.strip():
            chunks.append(current.strip())
        return chunks

    def _upsert_manifest_document(self, project_dir: Path, document: Dict[str, Any]) -> None:
        manifest_path = project_dir / "knowledge" / "manifest.json"
        manifest = self._read_json(manifest_path) or {"version": 1, "documents": []}
        knowledge_base = str(document.get("knowledge_base") or DEFAULT_KNOWLEDGE_BASE_NAME).strip() or DEFAULT_KNOWLEDGE_BASE_NAME
        document = {
            **document,
            "knowledge_base": knowledge_base,
            "kb_id": document.get("kb_id") or document.get("knowledge_base_id") or self._knowledge_base_id(knowledge_base),
        }
        documents = [
            item for item in manifest.get("documents", [])
            if item.get("doc_id") != document.get("doc_id")
        ]
        documents.append(document)
        manifest["version"] = manifest.get("version") or 1
        manifest["documents"] = documents
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    def _upsert_chunks(
        self,
        project_dir: Path,
        doc_id: str,
        file_name: str,
        relative_path: str,
        chunks: List[str],
        sender_identity: str = "",
    ) -> None:
        chunks_path = project_dir / "knowledge" / "chunks.jsonl"
        existing = []
        if chunks_path.exists():
            for line in chunks_path.read_text(encoding="utf-8").splitlines():
                if not line.strip():
                    continue
                try:
                    item = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if item.get("doc_id") != doc_id:
                    existing.append(item)

        for index, chunk in enumerate(chunks, 1):
            existing.append({
                "chunk_id": f"{doc_id}_{index:04d}",
                "doc_id": doc_id,
                "source_file_name": file_name,
                "relative_path": relative_path,
                "chunk_index": index,
                "sender_identity": sender_identity,
                "text": chunk,
            })

        chunks_path.write_text(
            "\n".join(json.dumps(item, ensure_ascii=False) for item in existing) + ("\n" if existing else ""),
            encoding="utf-8",
        )

    def select_relevant_documents(
        self,
        user_input: str,
        project_id: str = "",
        scene_id: str = "auto",
        llm_tools: Any = None,
        model_conf: Optional[Dict[str, Any]] = None,
        max_documents: int = 3,
        allowed_kb_ids: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        project_dir = self._resolve_project_dir(project_id)
        docs = self.list_knowledge_documents(project_dir.name)
        allowed = {str(item) for item in (allowed_kb_ids or []) if str(item or "").strip()}
        if allowed_kb_ids is not None:
            docs = [
                doc for doc in docs
                if str(doc.get("kb_id") or doc.get("knowledge_base_id") or self._knowledge_base_id(doc.get("knowledge_base") or "") or "") in allowed
            ]
        selected_ids: List[str] = []
        reason = ""

        if not docs:
            return {
                "documents": [],
                "document_ids": [],
                "context": "",
                "reason": "no_authorized_knowledge_base",
            }

        if llm_tools is not None and model_conf:
            try:
                selected_ids, reason = self._select_document_ids_with_ai(
                    user_input=user_input,
                    scene_id=scene_id,
                    docs=docs,
                    llm_tools=llm_tools,
                    model_conf=model_conf,
                    max_documents=max_documents,
                )
            except Exception:
                selected_ids = []
                reason = ""

        if not selected_ids:
            selected_ids = self._select_document_ids_by_rules(user_input, scene_id, docs, max_documents)
            reason = reason or "rule_fallback"

        selected_docs = self._read_selected_documents(project_dir, docs, selected_ids, max_documents)
        return {
            "documents": selected_docs,
            "document_ids": [doc["doc_id"] for doc in selected_docs],
            "context": self._format_document_context(selected_docs),
            "reason": reason,
        }

    def get_bundle(
        self,
        project_id: str = "",
        scene_id: str = "auto",
        user_input: str = "",
        document_context: str = "",
    ) -> ProjectBundle:
        project_dir = self._resolve_project_dir(project_id)
        project_meta = self._read_json(project_dir / "project.json")
        resolved_scene_id = scene_id or project_meta.get("default_scene") or "auto"
        if resolved_scene_id == "auto":
            resolved_scene_id = self.match_scene(user_input)

        scene_dir = project_dir / "scenes" / resolved_scene_id
        if not scene_dir.exists():
            resolved_scene_id = "general_presales"
            scene_dir = project_dir / "scenes" / resolved_scene_id

        scene_meta = self._read_json(scene_dir / "scene.json")
        global_prompt = self._read_text(project_dir / "global_prompt.md") or GLOBAL_OPERATOR_PROMPT
        scene_prompt = self._read_text(scene_dir / "prompt.md")
        materials_context = self._read_text(scene_dir / "materials.md")
        return ProjectBundle(
            project_id=project_dir.name,
            project_name=project_meta.get("name") or project_dir.name,
            scene_id=resolved_scene_id,
            scene_name=scene_meta.get("name") or resolved_scene_id,
            global_prompt=global_prompt,
            scene_prompt=scene_prompt,
            materials_context=materials_context,
            document_context=document_context,
        )

    def _select_document_ids_with_ai(
        self,
        user_input: str,
        scene_id: str,
        docs: List[Dict[str, Any]],
        llm_tools: Any,
        model_conf: Dict[str, Any],
        max_documents: int,
    ) -> tuple[List[str], str]:
        try:
            from aisec_agent.model.llm_typing import ProjectDocumentSelectionModel
        except ModuleNotFoundError as e:
            if e.name != "pydantic":
                raise
            ProjectDocumentSelectionModel = True

        catalog = [
            {
                "doc_id": doc.get("doc_id"),
                "title": doc.get("title"),
                "knowledge_base": doc.get("knowledge_base"),
                "domain": doc.get("domain"),
                "section": doc.get("section"),
                "category": doc.get("category"),
                "description": doc.get("description"),
                "keywords": doc.get("keywords", []),
            }
            for doc in docs
        ]
        prompt = f"""
你是项目资料路由器。请根据用户最新输入和当前场景，从文档清单里选择需要读取的资料文档。

规则：
- 最多选择 {max_documents} 个文档。
- 如果用户询问公司是谁、业务范围、平台介绍、综合能力，选择 company_positioning。
- 如果用户涉及关节、骨积液、软骨、干细胞、PRP、康养、医疗产品，选择 healthcare_medical。
- 如果用户涉及 AI 内容生产、自动发布、评论采集、录音纪要、智能硬件、私信引流、舆论管控，选择 ai_technology_hardware。
- 如果用户涉及跨境电商、AI选品、内容托管、留学移民、资源对接，选择 cross_border_services。
- 如果用户涉及招聘、岗位、候选人、简历、面试、薪资、HR、Boss 直聘，优先选择“招聘知识库”里的相关文档。
- 如果判断不明确，优先选择 company_positioning，再选择最接近的业务板块。

当前场景：{scene_id}
文档清单：
{json.dumps(catalog, ensure_ascii=False, indent=2)}

只返回 JSON，字段固定为：document_ids, reason。
""".strip()
        chat = getattr(llm_tools, model_conf["func_name"])
        response = chat(
            url=model_conf["url"],
            api_key=model_conf.get("key", ""),
            prompt=prompt,
            message=user_input,
            model=model_conf["model_name"],
            json_format=ProjectDocumentSelectionModel,
            stream=False,
            max_len_input=model_conf.get("max_len_input", 16000),
        )
        data = self._parse_json_response(response)
        allowed = {doc.get("doc_id") for doc in docs}
        selected_ids = [
            str(doc_id)
            for doc_id in data.get("document_ids", [])
            if str(doc_id) in allowed
        ]
        return selected_ids[:max_documents], str(data.get("reason", ""))

    @staticmethod
    def _boost_mock_business_doc_score(text: str, doc_id: str) -> int:
        groups = {
            "mock_sleep_assessment": ["睡眠", "睡不", "失眠", "入睡", "早醒", "压力", "熬夜", "精神不好"],
            "mock_joint_assessment": ["膝", "膝骨关节", "关节", "上下楼", "骨积液", "软骨", "软骨磨损", "老人", "疼", "肿", "干细胞", "prp", "注射", "门诊", "评估"],
            "mock_ai_dm_system": ["自动回复", "私信", "评论采集", "抖音评论", "视频摘要", "知识库", "模型", "引流系统"],
            "mock_cross_border_operation": ["跨境", "电商", "货源", "开店", "店铺", "托管", "选品", "海外"],
            "mock_conversion_hooks": ["钩子", "甜头", "资料包", "评估", "体验", "名额", "草料码", "微信", "入口", "领取"],
            "mock_health_compliance": ["医疗", "健康", "疗效", "治", "诊断", "睡眠", "膝", "关节", "活动", "合规"],
            "mock_comment_intent_samples": ["评论", "意向", "怎么回", "回复方向", "真的假的", "有用吗", "多少钱"],
            "mock_video_direction_samples": ["视频", "概述", "总结", "转写", "音频", "方向"],
        }
        return sum(4 for keyword in groups.get(doc_id, []) if keyword.lower() in text)

    @staticmethod
    def _document_priority_bonus(doc_id: str, text: str) -> int:
        recruitment_intent = re.search(
            r"招聘|岗位|候选人|简历|面试|薪资|福利|hr|boss|求职|入职|到岗|工程师|项目经理",
            text,
            flags=re.I,
        )
        if doc_id.startswith("recruitment_") and not recruitment_intent:
            return -30
        if doc_id.startswith("recruitment_") and re.search(r"招聘|岗位|候选人|简历|面试|薪资|福利|hr|boss|求职|入职|到岗|工程师|项目经理", text, flags=re.I):
            return 18
        if doc_id == "recruitment_ai_project_manager" and re.search(r"ai项目经理|项目经理|产品|需求|交付|推进", text, flags=re.I):
            return 22
        if doc_id == "recruitment_engineer_role" and re.search(r"工程师|后端|前端|开发|爬虫|接口|agent|技术", text, flags=re.I):
            return 22
        if doc_id == "recruitment_interview_process" and re.search(r"面试|简历|投递|到岗|流程|方便沟通", text, flags=re.I):
            return 20
        if doc_id == "recruitment_compensation_compliance" and re.search(r"薪资|工资|待遇|福利|社保|真假|录用|承诺", text, flags=re.I):
            return 20
        if doc_id in {"mock_video_direction_samples", "mock_comment_intent_samples"}:
            return -6
        if doc_id == "mock_joint_assessment" and re.search(r"膝骨关节|膝盖|上下楼|骨积液|软骨|干细胞|prp|注射|门诊评估", text, flags=re.I):
            return 24
        if doc_id == "mock_health_compliance" and re.search(r"医疗|膝骨关节|膝盖|骨积液|软骨|干细胞|prp|注射|疗效|治|诊断", text, flags=re.I):
            return 12
        business_keywords = {
            "mock_sleep_assessment": ["睡眠", "睡不", "失眠", "入睡", "早醒", "压力", "熬夜"],
            "mock_joint_assessment": ["膝", "膝骨关节", "关节", "上下楼", "骨积液", "软骨", "软骨磨损", "老人", "疼", "肿", "干细胞", "prp", "注射", "门诊", "评估"],
            "mock_ai_dm_system": ["自动回复", "私信", "评论采集", "抖音评论", "视频摘要", "知识库", "模型", "引流系统"],
            "mock_cross_border_operation": ["跨境", "电商", "货源", "开店", "店铺", "托管", "选品", "海外"],
        }
        return 8 if any(keyword.lower() in text for keyword in business_keywords.get(doc_id, [])) else 0

    @staticmethod
    def _score_document_values(text: str, values: List[Any]) -> int:
        score = 0
        for value in values:
            item = str(value or "").strip()
            if item and item.lower() in text:
                score += 2
            for token in re.split(r"[\s,，、/；;|。.!！?？：:\-]+", item):
                token = token.strip()
                if len(token) >= 2 and token.lower() in text:
                    score += 1
        return score

    @staticmethod
    def _select_document_ids_by_rules(
        user_input: str,
        scene_id: str,
        docs: List[Dict[str, Any]],
        max_documents: int,
    ) -> List[str]:
        text = (user_input or "").lower()
        scores: Dict[str, int] = {}
        scene_defaults = {
            "health_presales": "healthcare_medical",
            "ai_hardware_presales": "ai_technology_hardware",
            "cross_border_presales": "cross_border_services",
            "general_presales": "company_positioning",
            "auto": "company_positioning",
        }

        for doc in docs:
            doc_id = doc.get("doc_id", "")
            score = ProjectMaterialStore._score_document_values(
                text,
                [
                    doc.get("title"),
                    doc.get("knowledge_base"),
                    doc.get("domain"),
                    doc.get("section"),
                    doc.get("description"),
                    doc.get("category"),
                    *(doc.get("keywords", []) or []),
                ],
            )
            score += ProjectMaterialStore._boost_mock_business_doc_score(text, doc_id)
            score += ProjectMaterialStore._document_priority_bonus(doc_id, text)
            if doc_id == scene_defaults.get(scene_id):
                score += 1
            if score:
                scores[doc_id] = score

        if not scores:
            scores["company_positioning"] = 1
            if scene_defaults.get(scene_id) and scene_defaults[scene_id] != "company_positioning":
                scores[scene_defaults[scene_id]] = 1

        ordered = sorted(scores.items(), key=lambda item: (-item[1], item[0]))
        selected = [doc_id for doc_id, _ in ordered[:max_documents]]
        if "company_positioning" not in selected and len(selected) < max_documents:
            selected.insert(0, "company_positioning")
        return selected[:max_documents]

    def _read_selected_documents(
        self,
        project_dir: Path,
        docs: List[Dict[str, Any]],
        selected_ids: List[str],
        max_documents: int,
    ) -> List[Dict[str, Any]]:
        docs_by_id = {doc.get("doc_id"): doc for doc in docs}
        selected = []
        for doc_id in selected_ids[:max_documents]:
            doc = docs_by_id.get(doc_id)
            if not doc:
                continue
            content, resolved_relative = self._read_knowledge_document_text(project_dir, doc)
            selected.append({
                **doc,
                "sender_identity": self._display_sender_identity(doc.get("sender_identity")) or "品牌客服",
                "relative_path": resolved_relative,
                "content": content,
            })
        return selected

    @classmethod
    def _knowledge_document_path_candidates(cls, doc: Dict[str, Any]) -> List[str]:
        relative = str(doc.get("relative_path") or "").replace("\\", "/").strip().lstrip("/")
        knowledge_base = str(doc.get("knowledge_base") or DEFAULT_KNOWLEDGE_BASE_NAME).strip()
        domain = str(doc.get("domain") or doc.get("category") or "").strip()
        section = str(doc.get("section") or "").strip()
        title = Path(relative).name
        candidates: List[str] = []

        def add(path: str) -> None:
            normalized = str(path or "").replace("\\", "/").strip().lstrip("/")
            if normalized and normalized not in candidates:
                candidates.append(normalized)

        add(relative)
        prefix = "knowledge/files/"
        if relative.startswith(prefix) and knowledge_base:
            rest = relative[len(prefix):]
            kb_prefix = f"{prefix}{knowledge_base}/"
            if not relative.startswith(kb_prefix):
                add(kb_prefix + rest)
            else:
                rest_after_base = relative[len(kb_prefix):]
                parts = [part for part in rest_after_base.split("/") if part]
                if len(parts) >= 3:
                    add(kb_prefix + "/".join([parts[0], *parts[2:]]))
        if knowledge_base and domain and title:
            if section:
                add(f"{prefix}{knowledge_base}/{domain}/{section}/{title}")
            add(f"{prefix}{knowledge_base}/{domain}/{title}")
        return candidates

    def _read_knowledge_document_text(self, project_dir: Path, doc: Dict[str, Any]) -> tuple[str, str]:
        fallback_relative = str(doc.get("relative_path") or "").replace("\\", "/").strip().lstrip("/")
        for relative in self._knowledge_document_path_candidates(doc):
            path = project_dir / relative
            if path.is_file():
                return self._sanitize_prompt_document_text(self._read_text(path)), relative
        return "", fallback_relative

    def _manifest_document_file_exists(self, project_dir: Path, doc: Dict[str, Any]) -> bool:
        for relative in self._knowledge_document_path_candidates(doc):
            if (project_dir / relative).is_file():
                return True
        return False

    @staticmethod
    def _sanitize_prompt_document_text(content: str) -> str:
        lines = []
        for raw_line in str(content or "").splitlines():
            line = raw_line.strip()
            if not line:
                continue
            if re.search(r"(sender_identity|persona_identity|role_identity)\s*[:：]", line, flags=re.I):
                continue
            if re.search(r"(建议(私信)?身份|身份来源)", line):
                continue
            lines.append(re.sub(r"(?i)xx(?=[\u4e00-\u9fffA-Za-z0-9])", "", raw_line.rstrip()))
        return "\n".join(lines).strip()

    @staticmethod
    def _format_document_context(documents: List[Dict[str, Any]]) -> str:
        blocks = []
        for index, doc in enumerate(documents, 1):
            blocks.append(
                "\n".join([
                    f"[{index}] document_id: {doc.get('doc_id', '')}",
                    f"title: {doc.get('title', '')}",
                    f"knowledge_base: {doc.get('knowledge_base', '')}",
                    f"domain: {doc.get('domain', '')}",
                    f"section: {doc.get('section', '')}",
                    f"path: {doc.get('relative_path', '')}",
                    f"content:\n{ProjectMaterialStore._sanitize_prompt_document_text(doc.get('content', ''))}",
                ])
            )
        return "\n\n".join(blocks)

    @staticmethod
    def _parse_json_response(response: Any) -> Dict[str, Any]:
        if isinstance(response, dict):
            return response
        text = response.decode("utf-8", errors="replace") if isinstance(response, bytes) else str(response)
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            matches = re.findall(r"\{.*\}", text, flags=re.DOTALL)
            if matches:
                return json.loads(matches[0])
            raise

    @staticmethod
    def match_scene(user_input: str) -> str:
        text = user_input or ""
        scene_patterns = [
            ("health_presales", r"膝|关节|骨|积液|软骨|疼|痛|干细胞|PRP|康养|理疗|养老|胶原"),
            ("ai_hardware_presales", r"AI|视频|剪辑|发布|评论|数据|录音|纪要|硬件|爬虫|Agent|私信|舆论|小主机"),
            ("cross_border_presales", r"跨境|电商|选品|托管|留学|移民|海外|资源对接"),
        ]
        for scene_id, pattern in scene_patterns:
            if re.search(pattern, text, flags=re.IGNORECASE):
                return scene_id
        return "general_presales"

    def _project_dirs(self) -> List[Path]:
        return sorted([path for path in self.base_dir.iterdir() if path.is_dir()], key=lambda path: path.name)

    def _list_scenes(self, project_dir: Path) -> List[Dict[str, str]]:
        scenes = []
        for scene_dir in sorted((project_dir / "scenes").glob("*")):
            if not scene_dir.is_dir():
                continue
            meta = self._read_json(scene_dir / "scene.json")
            scenes.append(
                {
                    "scene_id": scene_dir.name,
                    "name": meta.get("name") or scene_dir.name,
                }
            )
        return scenes

    def _resolve_project_dir(self, project_id: str = "") -> Path:
        if project_id:
            candidate = self.base_dir / project_id
            if candidate.exists() and candidate.is_dir():
                return candidate
        return self.base_dir / self.default_project_id

    @staticmethod
    def _generated_project_id(name: str) -> str:
        return "project_" + uuid.uuid5(uuid.NAMESPACE_URL, name).hex[:12]

    @staticmethod
    def _read_json(path: Path) -> Dict[str, Any]:
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError):
            return {}

    @staticmethod
    def _read_text(path: Path) -> str:
        try:
            return path.read_text(encoding="utf-8").strip()
        except (FileNotFoundError, OSError):
            return ""
